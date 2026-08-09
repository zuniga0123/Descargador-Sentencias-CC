"""Descarga y parsing del botón 'Exportar resultados en EXCEL' del buscador.

Igual que en parsing.py: no fue posible confirmar contra el sitio real las
columnas exactas que trae el Excel exportado. El mapeo de columnas es
dinámico (por nombre de encabezado normalizado) para tolerar variaciones,
pero debe revisarse la primera vez que se corra contra el sitio real
(usar `inspeccionar-resultados` o abrir el .xlsx descargado manualmente).
"""

from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

from .http_client import RateLimitedClient
from .models import SearchResult
from .parsing import find_export_excel_link, split_tema_subtema
from .utils import normalize_key

logger = logging.getLogger(__name__)

# canonical_field -> alias de encabezado normalizados (sin acentos/espacios/mayúsculas) que se aceptan
HEADER_ALIASES: dict[str, list[str]] = {
    "radicado": ["radicado", "providencia", "numero", "numeroproviencia", "numeroprovidencia"],
    "url_ficha": ["enlace", "url", "link"],
    "fecha_sentencia": ["fechasentencia", "fechaprovidencia", "fecha"],
    "fecha_publicacion": ["fechapublicacion", "fechadepublicacion"],
    "sala": ["sala"],
    "ponente": ["ponente", "magistradoponente", "mp"],
    "tema_oficial": ["tema", "temas"],
    "subtema_oficial": ["subtema", "subtemas"],
    "resumen": ["resumen", "sintesis"],
}


def download_excel(client: RateLimitedClient, results_html: str, base_url: str) -> bytes | None:
    """Busca el link de exportación en la página de resultados y descarga el .xlsx. None si no se encontró."""
    link = find_export_excel_link(results_html, base_url)
    if not link:
        logger.info("No se encontró link de 'Exportar a Excel' en la página de resultados.")
        return None
    logger.info("Descargando exportación Excel: %s", link)
    return client.get_bytes(link)


def _map_headers(header_row: tuple) -> dict[int, str]:
    columnas: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        key = normalize_key(str(cell))
        for canonical, alias in HEADER_ALIASES.items():
            if key in alias and canonical not in columnas.values():
                columnas[idx] = canonical
                break
    return columnas


def parse_excel(data: bytes) -> list[SearchResult]:
    """Parsea el .xlsx exportado por el buscador a una lista de SearchResult.

    Lanza ValueError si el archivo no parece un .xlsx válido (por ejemplo, si
    el servidor devolvió una página de error HTML en vez del archivo).
    """
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - queremos degradar a fallback HTML
        raise ValueError(f"El contenido descargado no es un .xlsx válido: {exc}") from exc

    sheet = workbook.active
    filas = sheet.iter_rows(values_only=True)
    try:
        header = next(filas)
    except StopIteration:
        return []

    columnas = _map_headers(header)
    if "radicado" not in columnas.values():
        raise ValueError(
            "No se pudo identificar la columna de radicado en el Excel exportado; "
            "revisar HEADER_ALIASES en excel_export.py contra el archivo real."
        )

    resultados: list[SearchResult] = []
    for fila in filas:
        valores: dict[str, str] = {}
        for idx, campo in columnas.items():
            if idx < len(fila) and fila[idx] is not None:
                valores[campo] = str(fila[idx]).strip()

        radicado = valores.get("radicado")
        if not radicado:
            continue

        tema = valores.get("tema_oficial")
        subtema = valores.get("subtema_oficial")
        if tema and not subtema and "-" in tema:
            tema, subtema = split_tema_subtema(tema)

        resultados.append(
            SearchResult(
                radicado=radicado.upper(),
                url_ficha=valores.get("url_ficha", ""),
                fecha_sentencia=valores.get("fecha_sentencia"),
                fecha_publicacion=valores.get("fecha_publicacion"),
                sala=valores.get("sala"),
                ponente=valores.get("ponente"),
                tema_oficial=tema,
                subtema_oficial=subtema,
                resumen=valores.get("resumen"),
            )
        )

    return resultados
